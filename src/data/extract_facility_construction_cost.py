from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def load_costs_workbook(path: Path):
    wb = load_workbook(path, read_only=True)

    return wb.active


def get_costs(ws: Worksheet):
    result = []
    section_title = ""
    value = ""

    for linha in ws.iter_rows():
        if len(linha) < 4:
            continue

        if linha[0].value and str(linha[0].value).__contains__("FNDE"):
            section_title = linha[0].value
            continue
        elif linha[4].value and str(linha[4].value).startswith("VALOR:"):
            value = str(linha[6].value)

            # print(f"seção: {section_title}\nvalor: {value}")

            if not section_title or not value:
                raise ValueError(
                    "não é possível anexar um valor total sem sua seção e vice-versa"
                )

            result.append((section_title, value))

    return result
